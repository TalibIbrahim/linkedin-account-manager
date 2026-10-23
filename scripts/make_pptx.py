#!/usr/bin/env python3
"""Build an editable .pptx from a JSON layout spec.

Fallback renderer for when /mnt/skills/public/pptx/SKILL.md is not available
in the environment. Produces real text boxes/shapes (not a flattened image)
so the user can still edit everything in PowerPoint or re-upload to Canva.

Usage:
    python make_pptx.py spec.json output.pptx

Spec format (all pixel values assume 96px/inch, matching common social
image sizes like 1080x1350 or 1584x396):

{
  "slide_width_px": 1080,
  "slide_height_px": 1350,
  "background_color": "#FFFFFF",
  "slides": [
    {
      "shapes": [
        {
          "type": "rect",
          "x_px": 0, "y_px": 0, "w_px": 1080, "h_px": 12,
          "fill_color": "#2563EB"
        },
        {
          "type": "text",
          "x_px": 60, "y_px": 500, "w_px": 960, "h_px": 300,
          "text": "Shipping fast breaks things.\nI shipped the fix anyway.",
          "font_size": 54,
          "bold": true,
          "color": "#111111",
          "align": "left",
          "font": "Arial",
          "vertical_anchor": "middle"
        }
      ]
    }
  ]
}

One JSON file = one deck. A single-slide post image has one entry in
"slides"; a case-study carousel has 6-8.
"""

import json
import sys

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

PX_TO_EMU = 9525  # 1px at 96dpi, in EMU

ALIGN_MAP = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
}

ANCHOR_MAP = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}


def px(value):
    return Emu(int(round(value * PX_TO_EMU)))


def hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip("#")
    return RGBColor.from_string(hex_color.upper())


def add_text_box(slide, shape_spec):
    box = slide.shapes.add_textbox(
        px(shape_spec["x_px"]),
        px(shape_spec["y_px"]),
        px(shape_spec["w_px"]),
        px(shape_spec["h_px"]),
    )
    tf = box.text_frame
    tf.word_wrap = True
    anchor = shape_spec.get("vertical_anchor", "top")
    tf.vertical_anchor = ANCHOR_MAP.get(anchor, MSO_ANCHOR.TOP)

    lines = shape_spec["text"].split("\n")
    for i, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        paragraph.alignment = ALIGN_MAP.get(shape_spec.get("align", "left"), PP_ALIGN.LEFT)
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(shape_spec.get("font_size", 32))
        run.font.bold = bool(shape_spec.get("bold", False))
        run.font.name = shape_spec.get("font", "Arial")
        run.font.color.rgb = hex_to_rgb(shape_spec.get("color", "#111111"))


AUTOSHAPE_TYPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
}


def add_autoshape(slide, shape_spec):
    shape_type = AUTOSHAPE_TYPES[shape_spec["type"]]
    shape = slide.shapes.add_shape(
        shape_type,
        px(shape_spec["x_px"]),
        px(shape_spec["y_px"]),
        px(shape_spec["w_px"]),
        px(shape_spec["h_px"]),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_to_rgb(shape_spec.get("fill_color", "#000000"))
    if shape_spec.get("line_color"):
        shape.line.color.rgb = hex_to_rgb(shape_spec["line_color"])
        shape.line.width = Pt(shape_spec.get("line_width_pt", 1))
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False


SHAPE_BUILDERS = {
    "text": add_text_box,
    "rect": add_autoshape,
    "rounded_rect": add_autoshape,
    "oval": add_autoshape,
}


def build(spec):
    prs = Presentation()
    prs.slide_width = px(spec["slide_width_px"])
    prs.slide_height = px(spec["slide_height_px"])
    blank_layout = prs.slide_layouts[6]

    for slide_spec in spec["slides"]:
        slide = prs.slides.add_slide(blank_layout)

        bg = spec.get("background_color") or slide_spec.get("background_color")
        if bg:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = hex_to_rgb(bg)

        for shape_spec in slide_spec["shapes"]:
            builder = SHAPE_BUILDERS.get(shape_spec["type"])
            if builder is None:
                raise ValueError(f"Unknown shape type: {shape_spec['type']}")
            builder(slide, shape_spec)

    return prs


def main():
    if len(sys.argv) != 3:
        print("Usage: python make_pptx.py spec.json output.pptx", file=sys.stderr)
        sys.exit(1)

    spec_path, out_path = sys.argv[1], sys.argv[2]
    with open(spec_path, "r", encoding="utf-8") as f:
        spec = json.load(f)

    prs = build(spec)
    prs.save(out_path)
    print(f"Wrote {out_path} ({len(spec['slides'])} slide(s), "
          f"{spec['slide_width_px']}x{spec['slide_height_px']}px)")


if __name__ == "__main__":
    main()
